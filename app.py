from flask import Flask, request, jsonify, send_file, send_from_directory
from flask_cors import CORS
import sqlite3
import os
from io import BytesIO
from openpyxl import Workbook
from datetime import datetime

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.path.join(BASE_DIR, 'data.db')

app = Flask(__name__, static_folder='public', static_url_path='/')
CORS(app)

# Ensure DB and tables
def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn

def init_db():
    conn = get_db()
    cur = conn.cursor()
    # Create customers table if not exists
    cur.execute('''CREATE TABLE IF NOT EXISTS customers (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        name TEXT NOT NULL
    )''')
    # Create sales table with new columns (invoice_number, product_code)
    cur.execute('''CREATE TABLE IF NOT EXISTS sales (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        customer_id INTEGER NOT NULL,
        type TEXT,
        date TEXT,
        amount REAL DEFAULT 0,
        qty INTEGER DEFAULT 0,
        notes TEXT,
        invoice_number TEXT DEFAULT '',
        product_code TEXT DEFAULT '',
        FOREIGN KEY(customer_id) REFERENCES customers(id)
    )''')
    # In case the table existed before without the new columns, add them
    cur.execute("PRAGMA table_info(sales)")
    cols = [r[1] for r in cur.fetchall()]
    if 'invoice_number' not in cols:
        cur.execute("ALTER TABLE sales ADD COLUMN invoice_number TEXT DEFAULT ''")
    if 'product_code' not in cols:
        cur.execute("ALTER TABLE sales ADD COLUMN product_code TEXT DEFAULT ''")

    conn.commit()
    conn.close()

init_db()

# Serve the frontend
@app.route('/')
def index():
    return send_from_directory(app.static_folder, 'index.html')

# API: add customer
@app.route('/api/customers', methods=['POST'])
def add_customer():
    data = request.get_json() or {}
    name = data.get('name')
    if not name:
        return jsonify({'error': 'name required'}), 400
    conn = get_db()
    cur = conn.cursor()
    cur.execute('INSERT INTO customers(name) VALUES(?)', (name,))
    conn.commit()
    cid = cur.lastrowid
    conn.close()
    return jsonify({'id': cid, 'name': name})

# API: list customers
@app.route('/api/customers', methods=['GET'])
def list_customers():
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT * FROM customers ORDER BY name')
    rows = [dict(r) for r in cur.fetchall()]
    conn.close()
    return jsonify(rows)

# Helper to generate invoice number based on date and sequence for that day
def generate_invoice_number_for_date(conn, date_str):
    # date_str expected in YYYY-MM-DD
    cur = conn.cursor()
    cur.execute('SELECT COUNT(*) as cnt FROM sales WHERE date = ?', (date_str,))
    row = cur.fetchone()
    count = row['cnt'] if row else 0
    seq = count + 1
    datepart = date_str.replace('-', '')
    return f'INV-{datepart}-{seq:04d}'

# API: add sale (invoice_number now auto-generated)
@app.route('/api/sales', methods=['POST'])
def add_sale():
    data = request.get_json() or {}
    customer_id = data.get('customer_id')
    date = data.get('date')
    if not customer_id or not date:
        return jsonify({'error': 'customer_id and date required'}), 400
    type_ = data.get('type', '')
    amount = data.get('amount', 0) or 0
    qty = data.get('qty', 0) or 0
    notes = data.get('notes', '')
    product_code = data.get('product_code', '')

    conn = get_db()
    cur = conn.cursor()
    # Generate invoice number based on date and existing count for that date
    invoice_number = generate_invoice_number_for_date(conn, date)
    cur.execute('INSERT INTO sales(customer_id, type, date, amount, qty, notes, invoice_number, product_code) VALUES(?,?,?,?,?,?,?,?)',
                (customer_id, type_, date, amount, qty, notes, invoice_number, product_code))
    conn.commit()
    sid = cur.lastrowid
    conn.close()
    return jsonify({'id': sid, 'invoice_number': invoice_number})

# API: delete sale
@app.route('/api/sales/<int:sale_id>', methods=['DELETE'])
def delete_sale(sale_id):
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT * FROM sales WHERE id = ?', (sale_id,))
    row = cur.fetchone()
    if not row:
        conn.close()
        return jsonify({'error': 'not found'}), 404
    cur.execute('DELETE FROM sales WHERE id = ?', (sale_id,))
    conn.commit()
    conn.close()
    return jsonify({'deleted': sale_id})

# API: list sales (optionally by customer)
@app.route('/api/sales', methods=['GET'])
def list_sales():
    customer_id = request.args.get('customer_id')
    conn = get_db()
    cur = conn.cursor()
    if customer_id:
        cur.execute('SELECT s.*, c.name as customer_name FROM sales s JOIN customers c ON s.customer_id=c.id WHERE customer_id=? ORDER BY date DESC, id DESC', (customer_id,))
    else:
        cur.execute('SELECT s.*, c.name as customer_name FROM sales s JOIN customers c ON s.customer_id=c.id ORDER BY date DESC, id DESC')
    rows = [dict(r) for r in cur.fetchall()]
    conn.close()
    return jsonify(rows)

# Export single customer (includes new fields)
@app.route('/export/customer/<int:customer_id>', methods=['GET'])
def export_customer(customer_id):
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT * FROM customers WHERE id = ?', (customer_id,))
    c = cur.fetchone()
    if not c:
        conn.close()
        return 'Customer not found', 404
    cur.execute('SELECT * FROM sales WHERE customer_id = ? ORDER BY date', (customer_id,))
    sales = [dict(r) for r in cur.fetchall()]
    conn.close()

    wb = Workbook()
    # Summary sheet
    summary = wb.active
    summary.title = 'ملخص'
    summary.append(['الاسم', 'عدد العمليات', 'إجمالي المبيعات', 'إجمالي الكمية'])
    total_amount = sum((s.get('amount') or 0) for s in sales)
    total_qty = sum((s.get('qty') or 0) for s in sales)
    summary.append([c['name'], len(sales), total_amount, total_qty])

    # Operations sheet (with invoice and product)
    ws = wb.create_sheet('العمليات')
    ws.append(['المعرف', 'التاريخ', 'النوع', 'الكمية', 'المبلغ', 'رقم الفاتورة', 'رمز المنتج', 'ملاحظات'])
    for s in sales:
        ws.append([s['id'], s['date'], s.get('type', ''), s.get('qty', 0), s.get('amount', 0), s.get('invoice_number', ''), s.get('product_code', ''), s.get('notes', '')])

    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)
    filename = f'customer_{customer_id}.xlsx'
    return send_file(bio, as_attachment=True, download_name=filename, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

# Export all customers (includes new fields)
@app.route('/export/all', methods=['GET'])
def export_all():
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT * FROM customers ORDER BY name')
    customers = [dict(r) for r in cur.fetchall()]

    wb = Workbook()
    overview = wb.active
    overview.title = 'ملخص الكل'
    overview.append(['المعرف', 'العميل', 'عدد العمليات', 'إجمالي المبيعات', 'إجمالي الكمية'])

    for c in customers:
        cur.execute('SELECT * FROM sales WHERE customer_id = ? ORDER BY date', (c['id'],))
        sales = [dict(r) for r in cur.fetchall()]
        total = sum((s.get('amount') or 0) for s in sales)
        qty = sum((s.get('qty') or 0) for s in sales)
        overview.append([c['id'], c['name'], len(sales), total, qty])

        sh = wb.create_sheet(f'عميل_{c["id"]}')
        sh.append(['المعرف', 'التاريخ', 'النوع', 'الكمية', 'المبلغ', 'رقم الفاتورة', 'رمز المنتج', 'ملاحظات'])
        for s in sales:
            sh.append([s['id'], s['date'], s.get('type', ''), s.get('qty', 0), s.get('amount', 0), s.get('invoice_number', ''), s.get('product_code', ''), s.get('notes', '')])

    conn.close()
    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)
    filename = 'all_customers.xlsx'
    return send_file(bio, as_attachment=True, download_name=filename, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

# Export invoices reference file
@app.route('/export/invoices', methods=['GET'])
def export_invoices():
    conn = get_db()
    cur = conn.cursor()
    cur.execute('SELECT s.id, s.date, s.invoice_number, c.name as customer_name, s.amount, s.product_code FROM sales s JOIN customers c ON s.customer_id=c.id ORDER BY s.date, s.id')
    rows = [dict(r) for r in cur.fetchall()]
    conn.close()

    wb = Workbook()
    ws = wb.active
    ws.title = 'رقم العمليات'
    ws.append(['المعرف', 'التاريخ', 'رقم الفاتورة', 'العميل', 'المبلغ', 'رمز المنتج'])
    for r in rows:
        ws.append([r.get('id'), r.get('date'), r.get('invoice_number'), r.get('customer_name'), r.get('amount'), r.get('product_code')])

    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)
    filename = 'invoices_reference.xlsx'
    return send_file(bio, as_attachment=True, download_name=filename, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

if __name__ == '__main__':
    port = int(os.environ.get('PORT', 3000))
    print(f"Starting Flask on http://localhost:{port}")
    app.run(host='0.0.0.0', port=port)
